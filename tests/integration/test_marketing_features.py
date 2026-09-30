"""Export, digest, trends, review queue and visibility of undated promotions."""
import csv
import io
import uuid
from datetime import datetime, timedelta, timezone

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import text

from app.db.session import SessionLocal
from app.main import app
from app.models.auth import AuditLog, User
from app.models.entity import Brand, Competitor, Retailer
from app.models.promotion import Promotion, PromotionEvidence
from app.models.promotion_change import PromotionChangeEvent
from app.models.resolution import ReviewQueue
from app.services import auth as auth_service

PASSWORD = "Correct-Horse-Battery-9"


@pytest.fixture()
def env():
    db = SessionLocal()
    try:
        db.execute(text("SELECT 1"))
    except Exception as exc:
        pytest.skip(f"PostgreSQL is not reachable: {exc}")
    auth_service.ip_throttle.clear()
    tag = uuid.uuid4().hex[:8]
    users = {}
    for role in ("VIEWER", "ANALYST"):
        name = f"mk_{role.lower()}_{tag}"
        auth_service.create_user(db, username=name, password=PASSWORD, role=role)
        users[role] = name
    db.commit()
    retailer = db.query(Retailer).filter(Retailer.name == "Indomaret").one()
    competitor = db.query(Competitor).first()
    brands = db.query(Brand).limit(2).all()
    now = datetime.now(timezone.utc)
    promos = []

    def promo(name, **kw):
        base = dict(product_name=f"Zz{tag} {name}", category="BISCUIT", promotion_type="DISCOUNT", discount_percentage=25.0,
                    retailer_id=retailer.id, competitor_id=competitor.id, status="ACTIVE", last_seen_at=now,
                    first_seen_at=now, rank_score=0.9)
        base.update(kw)
        p = Promotion(**base)
        db.add(p)
        db.flush()
        promos.append(p)
        return p

    def login(role):
        c = TestClient(app)
        r = c.post("/api/v1/auth/login", json={"username": users[role], "password": PASSWORD})
        assert r.status_code == 200
        return c, r.json()["csrf_token"]

    ctx = type("Ctx", (), dict(db=db, tag=tag, now=now, promo=staticmethod(promo), login=staticmethod(login),
                               brands=brands, competitor=competitor, users=users, promos=promos))
    yield ctx
    db.rollback()
    ids = [p.id for p in promos]
    if ids:
        db.query(ReviewQueue).filter(ReviewQueue.promotion_id.in_(ids)).delete(synchronize_session=False)
        db.query(PromotionChangeEvent).filter(PromotionChangeEvent.promotion_id.in_(ids)).delete(synchronize_session=False)
        db.query(PromotionEvidence).filter(PromotionEvidence.promotion_id.in_(ids)).delete(synchronize_session=False)
        db.query(Promotion).filter(Promotion.id.in_(ids)).delete(synchronize_session=False)
    for name in users.values():
        db.query(AuditLog).filter(AuditLog.username == name).delete(synchronize_session=False)
        db.query(User).filter(User.username == name).delete(synchronize_session=False)
    db.commit()
    db.close()


def test_undated_promotions_are_visible_and_expired_are_not(env):
    dated = env.promo("Dated", start_date=env.now - timedelta(days=1), end_date=env.now + timedelta(days=4))
    undated = env.promo("Undated", status="UNKNOWN", rank_score=0.8)
    stale = env.promo("Stale undated", status="UNKNOWN", last_seen_at=env.now - timedelta(days=40))
    expired = env.promo("Expired", end_date=env.now - timedelta(days=1), status="EXPIRED")
    env.db.commit()
    c, _ = env.login("VIEWER")
    items = c.get(f"/api/v1/promotions/top10?q=Zz{env.tag}").json()["promotions"]
    names = {i["product_name"]: i for i in items}
    assert f"Zz{env.tag} Dated" in names and f"Zz{env.tag} Undated" in names
    assert f"Zz{env.tag} Stale undated" not in names and f"Zz{env.tag} Expired" not in names
    assert names[f"Zz{env.tag} Undated"]["dates_stated"] is False
    assert names[f"Zz{env.tag} Undated"]["valid_until"] == "Dates not stated"
    assert names[f"Zz{env.tag} Dated"]["channel"] == "Modern Trade"


def test_export_csv_is_safe_and_role_protected(env):
    env.promo("x", product_name=f'=HYPERLINK("http://evil.test","click") Zz{env.tag}')
    p = env.promo("Normal biskuit")
    env.db.add(PromotionEvidence(promotion_id=p.id, evidence_text="+cmd|' /C calc'!A0", source_url="https://x.test/a"))
    env.db.commit()

    viewer, _ = env.login("VIEWER")
    assert viewer.get("/api/v1/promotions/export").status_code == 403

    analyst, _ = env.login("ANALYST")
    r = analyst.get(f"/api/v1/promotions/export?format=csv&q=Zz{env.tag}")
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    assert r.content.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert rows[0][:3] == ["Rank", "Product", "Brand"]
    products = [row[1] for row in rows[1:]]
    assert any(p.startswith("'=HYPERLINK") for p in products)           # formula neutralised
    assert not any(p.startswith("=") for p in products)
    assert any(row[18].startswith("'+cmd") for row in rows[1:])          # evidence column neutralised too

    x = analyst.get(f"/api/v1/promotions/export?format=xlsx&q=Zz{env.tag}")
    assert x.status_code == 200
    ws = load_workbook(io.BytesIO(x.content)).active
    assert ws["B1"].value == "Product" and ws.max_row == 3
    assert analyst.get("/api/v1/promotions/export?format=pdf").status_code == 422


def test_digest_lists_new_changed_and_ending(env):
    new = env.promo("Brand new")
    ending = env.promo("Ending soon", first_seen_at=env.now - timedelta(days=30),
                       start_date=env.now - timedelta(days=10), end_date=env.now + timedelta(days=2))
    env.db.add(PromotionChangeEvent(
        promotion_id=ending.id, event_type="PRICE_OR_VALUE_CHANGED", field_name="promo_price", previous_value=9000,
        new_value=7500, change_impact=0.9, observed_at=env.now, event_fingerprint=uuid.uuid4().hex + uuid.uuid4().hex[:32],
        created_at=env.now))
    env.db.commit()
    c, _ = env.login("VIEWER")
    d = c.get("/api/v1/promotions/digest?days=7").json()
    assert d["summary"]["new"] >= 1 and d["summary"]["ending_soon"] >= 1 and d["summary"]["changed"] >= 1
    assert f"Zz{env.tag} Brand new" in [p["product"] for p in d["new_promotions"]] or d["summary"]["new"] > 15
    assert f"Zz{env.tag} Ending soon" in [p["product"] for p in d["ending_soon"]]
    assert any(ch["product"] == f"Zz{env.tag} Ending soon" and ch["new"] == 7500 for ch in d["changes"])


def test_digest_email_is_escaped_and_plain_text_available():
    from app.services.digest import render_digest
    digest = {"summary": {"new": 1, "changed": 0, "ending_soon": 0}, "days": 7,
              "new_promotions": [{"product": "<script>alert(1)</script>", "competitor": "Mayora", "outlet": "Indomaret",
                                  "discount": 20, "promo_price": None, "valid_until": None}],
              "ending_soon": [], "changes": []}
    subject, text_body, html_body = render_digest(digest)
    assert "<script>" not in html_body and "&lt;script&gt;" in html_body
    assert "Mayora" in text_body and "1 new" in subject


def test_trends_shape(env):
    env.promo("Trend item")
    env.db.commit()
    c, _ = env.login("VIEWER")
    t = c.get("/api/v1/stats/trends?weeks=8").json()
    assert len(t["weeks"]) == 8
    mine = next(x for x in t["competitors"] if x["competitor"] == env.competitor.name)
    assert len(mine["counts"]) == 8 and mine["counts"][-1] >= 1
    assert c.get("/api/v1/stats/trends?weeks=1").status_code == 422


def test_review_queue_approve_reject_and_permissions(env):
    p = env.promo("Needs review")
    approve = ReviewQueue(entity_type="BRAND", entity_id=env.brands[0].id, candidate_entity_id=env.brands[1].id,
                          promotion_id=p.id, reason="Fuzzy match", confidence=0.7, priority=2, status="PENDING")
    no_candidate = ReviewQueue(entity_type="BRAND", promotion_id=p.id, reason="Unknown brand", confidence=0.3, status="PENDING")
    env.db.add_all([approve, no_candidate])
    env.db.commit()

    viewer, vcsrf = env.login("VIEWER")
    assert viewer.get("/api/v1/review/").status_code == 403

    c, csrf = env.login("ANALYST")
    hdr = {"X-CSRF-Token": csrf}
    listed = {i["id"]: i for i in c.get("/api/v1/review/").json()}
    assert listed[str(approve.id)]["can_approve"] is True and listed[str(no_candidate.id)]["can_approve"] is False

    assert c.post(f"/api/v1/review/{approve.id}/resolve", json={"decision": "APPROVED"}).status_code == 403  # CSRF
    assert c.post(f"/api/v1/review/{no_candidate.id}/resolve", json={"decision": "APPROVED"}, headers=hdr).status_code == 422
    assert c.post(f"/api/v1/review/{approve.id}/resolve", json={"decision": "MAYBE"}, headers=hdr).status_code == 422
    r = c.post(f"/api/v1/review/{approve.id}/resolve", json={"decision": "APPROVED", "notes": "looks right"}, headers=hdr)
    assert r.status_code == 200 and r.json()["status"] == "APPROVED"
    env.db.expire_all()
    assert env.db.get(Promotion, p.id).brand_id == env.brands[1].id            # promotion linked to confirmed brand
    assert c.post(f"/api/v1/review/{approve.id}/resolve", json={"decision": "REJECTED"}, headers=hdr).status_code == 409
    assert c.post(f"/api/v1/review/{no_candidate.id}/resolve", json={"decision": "REJECTED"}, headers=hdr).status_code == 200
