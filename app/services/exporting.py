"""CSV / Excel export helpers."""
from __future__ import annotations

import csv
import io
from typing import Any, Dict, List, Tuple

EXPORT_HEADERS = [
    "Rank", "Product", "Brand", "Competitor", "Category", "Pack size", "Outlet", "Channel", "Promotion type",
    "Regular price", "Promo price", "Discount %", "Valid from", "Valid until", "Dates stated", "Score",
    "AI confidence", "Source URL", "Evidence quote", "Last verified",
]
_KEYS = [
    "rank", "product", "brand", "competitor", "category", "pack_size", "outlet", "channel", "promotion_type",
    "regular_price", "promo_price", "discount", "valid_from", "valid_until", "dates_stated", "score",
    "confidence", "source_url", "evidence", "last_verified",
]
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def safe_cell(value: Any) -> Any:
    """Neutralise spreadsheet formula injection: scraped text must never run as a formula."""
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def build_export(rows: List[Dict[str, Any]], fmt: str) -> Tuple[bytes, str, str]:
    table = [[safe_cell(r.get(k)) for k in _KEYS] for r in rows]
    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb = Workbook()
        ws = wb.active
        ws.title = "Promotions"
        ws.append(EXPORT_HEADERS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in table:
            ws.append(row)
        ws.freeze_panes = "A2"
        for col, header in enumerate(EXPORT_HEADERS, start=1):
            ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = min(48, max(12, len(header) + 4))
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(EXPORT_HEADERS)
    writer.writerows(table)
    # UTF-8 BOM so Excel opens Indonesian characters correctly.
    return ("\ufeff" + buf.getvalue()).encode("utf-8"), "text/csv; charset=utf-8", "csv"
